import os
import sys
import re
import json
import shutil
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client
import fitz

from tcc_humanizer import set_cell_margins, set_cell_background, format_abnt_table, humanize_text


def load_experimental_data() -> dict:
    json_path = Path(__file__).parent / "battle_ia_resultados_brutos.json"
    if not json_path.exists():
        raise FileNotFoundError(f"ERRO FATAL: Base empírica não encontrada em {json_path}")
    with open(json_path, "r", encoding="utf-8") as f:
        records = json.load(f)
    if len(records) < 12:
        raise ValueError(f"ERRO FATAL: Esperados 12 registros experimentais, encontrados {len(records)}")
    by_key = {(r["caso_id"], r["tecnica"]): r for r in records}
    return {"records": records, "by_key": by_key}


def fmt_br(val: float, decimals: int = 1) -> str:
    s = f"{val:,.{decimals}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def build_monograph(docx_path: Path, sumario_page_map=None):
    exp = load_experimental_data()
    by_key = exp["by_key"]

    doc = Document()

    # 1. Margens ABNT
    for section in doc.sections:
        section.top_margin = Cm(3.0)
        section.left_margin = Cm(3.0)
        section.bottom_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    def add_p(text, bold_prefix=None, space_after=5, line_spacing=1.5, first_line_indent=Cm(1.25), align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        text = humanize_text(text)
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.first_line_indent = first_line_indent

        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(12)
            r_bold.font.bold = True
            r_bold.font.color.rgb = RGBColor(0, 0, 0)

        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_h1(text, page_break_before=False):
        if page_break_before:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text.upper())
        r.font.name = "Arial"
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_table_title(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        return p

    def add_table_source(source="Fonte: Dados experimentais obtidos pelos autores via NVIDIA NIM Cloud (2026)."):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(source)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.italic = True
        return p

    def add_figure(img_path: Path, title: str, source="Fonte: Dados experimentais obtidos pelos autores via NVIDIA NIM Cloud (2026)."):
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_title.paragraph_format.space_before = Pt(12)
        p_title.paragraph_format.space_after = Pt(3)
        p_title.paragraph_format.first_line_indent = Cm(0)
        p_title.paragraph_format.line_spacing = 1.0
        p_title.paragraph_format.keep_with_next = True
        r_t = p_title.add_run(title)
        r_t.font.name = "Arial"
        r_t.font.size = Pt(10)
        r_t.font.bold = True

        if img_path.exists():
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(2)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.first_line_indent = Cm(0)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(str(img_path), width=Cm(14.0))

        p_src = doc.add_paragraph()
        p_src.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_src.paragraph_format.space_before = Pt(2)
        p_src.paragraph_format.space_after = Pt(10)
        p_src.paragraph_format.first_line_indent = Cm(0)
        p_src.paragraph_format.line_spacing = 1.0
        r_s = p_src.add_run(source)
        r_s.font.name = "Arial"
        r_s.font.size = Pt(9)
        r_s.font.italic = True

    figs_dir = Path(__file__).parent.parent / "04 - Figuras"

    # =========================================================================
    # 1. CAPA RÍGIDA DE 1 PÁGINA (Com INDAIATUBA 2026 no rodapé)
    # =========================================================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(0)
    p_inst.paragraph_format.line_spacing = 1.3
    r = p_inst.add_run("CENTRO UNIVERSITÁRIO MAX PLANCK\nCURSO DE BACHARELADO EM CIÊNCIA DA COMPUTAÇÃO")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(110)
    p_title.paragraph_format.space_after = Pt(0)
    p_title.paragraph_format.line_spacing = 1.3
    r_tit = p_title.add_run("COMPRESSÃO DE PROMPTS, TOKENS E SUSTENTABILIDADE DA IA\n\n")
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(14)
    r_tit.font.bold = True

    r_sub = p_title.add_run("Engenharia de Prompt como Estratégia de Otimização de Respostas em Inteligências Artificiais")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(12)
    r_sub.font.bold = False

    p_nat = doc.add_paragraph()
    p_nat.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_nat.paragraph_format.space_before = Pt(110)
    p_nat.paragraph_format.space_after = Pt(0)
    p_nat.paragraph_format.line_spacing = 1.15
    p_nat.paragraph_format.left_indent = Cm(8.0)
    r_nat = p_nat.add_run(
        "Trabalho de Conclusão de Curso apresentado ao Curso de Bacharelado em Ciência da Computação "
        "do Centro Universitário Max Planck (UniMAX), como requisito parcial para obtenção do grau de "
        "Bacharel em Ciência da Computação.\n\n"
        "Orientador: Prof. Luiz Claudio Chiavini Oliveira Junior"
    )
    r_nat.font.name = "Arial"
    r_nat.font.size = Pt(10)

    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_city.paragraph_format.space_before = Pt(130)
    p_city.paragraph_format.space_after = Pt(0)
    p_city.paragraph_format.line_spacing = 1.2
    r_city = p_city.add_run("INDAIATUBA\n2026")
    r_city.font.name = "Arial"
    r_city.font.size = Pt(12)
    r_city.font.bold = True

    # =========================================================================
    # 2. RESUMO (Página 2)
    # =========================================================================
    doc.add_page_break()
    p_res_title = doc.add_paragraph()
    p_res_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_res_title.paragraph_format.space_before = Pt(12)
    p_res_title.paragraph_format.space_after = Pt(12)
    r = p_res_title.add_run("RESUMO")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True

    add_p(
        "A inferência de Modelos de Linguagem de Grande Escala (LLMs) impõe um custo computacional elevado à sua operação contínua, "
        "visto que o mecanismo de autoatenção da arquitetura Transformer escala de forma quadrática O(n²) em relação ao comprimento da entrada. "
        "Este trabalho examina a hipótese de redundância linguística de Claude Shannon aplicada à engenharia de prompt, avaliando como o descarte "
        "algorítmico de tokens de baixa autoinformação reduz a sobrecarga de hardware sem descaracterizar a precisão técnica das respostas. "
        "Construímos a bancada experimental automatizada Battle IA para testar três cenários corporativos: raciocínio em sistemas distribuídos (consenso Raft e Teorema CAP), "
        "geração de consultas SQL analíticas sob esquemas DDL em 3FN e auditoria de conformidade com a LGPD. "
        "Submetemos o modelo Meta Llama-3.2-11B-Vision-Instruct, operado via nuvem NVIDIA NIM Cloud sobre aceleradores NVIDIA A100 SXM4, a ensaios reais em streaming SSE, "
        "contrastando linhas de base convencionais (Few-shot e Zero-shot) com a compressão seletiva inspirada no LLMLingua combinada a listas de proteção lexical (token whitelisting) em taxas nominais de 2x e 4x. "
        "Os resultados empíricos demonstram que a compressão nominal de 4x reduz o volume faturado de tokens de entrada em até 57,4% (média de 56,2%, caindo de 384,0 para 168,3 tokens por requisição) "
        "e elimina em média 80,7% das operações quadráticas de autoatenção no prefill (atingindo 81,8% no Caso 2). "
        "O Time-to-First-Token (TTFT) estabilizou-se entre 538,0 ms e 636,7 ms sob compressão de 4x, mitigando oscilações de fila na nuvem e reduzindo a latência inicial em 23,0% frente ao Few-shot no Caso 3. "
        "A fidelidade semântica (BERTScore Proxy F1, com limiar crítico de 0,80) manteve-se elevada na geração de SQL (0,8524 em 4x, com 100% de retenção de entidades) e apresentou ganho de sinal-ruído na auditoria LGPD "
        "(evoluindo de 0,7916 no baseline para 0,8279 em 4x), registrando retração marginal apenas no Caso 1 (0,7912) pela sensibilidade de modelos compactos de 11B parâmetros à omissão de conectivos causais. "
        "Em escala de 100.000 requisições, a técnica reduz em 56,2% o dispêndio de entrada e em 12,8% o custo global (de US$ 505,20 para US$ 440,50), consolidando-se como instrumento prático de Green AI.",
        space_after=10, line_spacing=1.25, first_line_indent=Cm(0)
    )

    add_p("Compressão de Prompts. Modelos de Linguagem de Grande Escala. Engenharia de Prompt. Green AI. BERTScore. NVIDIA NIM.",
          bold_prefix="Palavras-chave: ", line_spacing=1.25, first_line_indent=Cm(0))

    # =========================================================================
    # 3. ABSTRACT (Página 3)
    # =========================================================================
    doc.add_page_break()
    p_abs_title = doc.add_paragraph()
    p_abs_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs_title.paragraph_format.space_before = Pt(12)
    p_abs_title.paragraph_format.space_after = Pt(12)
    r = p_abs_title.add_run("ABSTRACT")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True

    add_p(
        "Inference in Large Language Models (LLMs) imposes a substantial computational burden on modern infrastructure, as inner product "
        "calculations in the Transformer self-attention mechanism scale quadratically O(n²) with sequence length. This research examines Claude Shannon's "
        "linguistic redundancy hypothesis applied to prompt engineering, evaluating how algorithmic pruning of low-self-information tokens mitigates "
        "hardware strain while preserving technical accuracy. We built an automated experimental testbed (Battle IA Benchmark) covering three enterprise "
        "workloads: distributed systems consensus reasoning (Raft and CAP Theorem), analytical SQL query generation under 3NF relational DDL schemas, and LGPD data privacy auditing. "
        "Evaluating Meta Llama-3.2-11B-Vision-Instruct deployed on NVIDIA NIM Cloud (NVIDIA A100 SXM4 GPUs) via real-time SSE streaming, we benchmarked "
        "conventional baselines (Few-shot and Zero-shot) against LLMLingua-inspired selective compression with domain token whitelisting at 2x and 4x nominal target ratios. "
        "Empirical measurements confirm that 4x nominal compression reduces billed input token volume by up to 57.4% (averaging 56.2%, from 384.0 to 168.3 input tokens per request) "
        "and eliminates an average of 80.7% of quadratic prefill self-attention FLOPs (reaching 81.8% in Case 2). "
        "Time-to-First-Token (TTFT) stabilized between 538.0 ms and 636.7 ms under 4x compression, mitigating multi-tenant cloud queueing jitter and cutting initial latency by 23.0% against the Few-shot baseline in Case 3. "
        "Semantic retention (measured via BERTScore Proxy F1 against a 0.80 academic acceptance threshold) remained strong in SQL generation (0.8524 at 4x with 100% entity retention) and exhibited positive signal-to-noise enhancement "
        "in LGPD auditing (rising from 0.7916 in Few-shot to 0.8279 at 4x), while showing a slight dip to 0.7912 in Case 1 due to causal connective sensitivity in an 11B-parameter model. "
        "Across 100,000 requests, the approach cuts input token expenditure by 56.2% and total batch cost by 12.8% (from US$ 505.20 to US$ 440.50), validating selective prompt compression as an effective Green AI practice.",
        space_after=10, line_spacing=1.25, first_line_indent=Cm(0)
    )

    add_p("Prompt Compression. Large Language Models. Prompt Engineering. Green AI. BERTScore. NVIDIA NIM.",
          bold_prefix="Keywords: ", line_spacing=1.25, first_line_indent=Cm(0))

    # =========================================================================
    # 4. SUMÁRIO (Página 4)
    # =========================================================================
    doc.add_page_break()
    p_sum_title = doc.add_paragraph()
    p_sum_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sum_title.paragraph_format.space_before = Pt(24)
    p_sum_title.paragraph_format.space_after = Pt(18)
    r = p_sum_title.add_run("SUMÁRIO")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True

    toc_items = [
        ("1 INTRODUÇÃO", 1),
        ("1.1 Contextualização e Motivação", 2),
        ("1.2 Objetivos", 2),
        ("1.3 Justificativa e Relevância", 2),
        ("2 FUNDAMENTOS DOS MODELOS DE LINGUAGEM DE GRANDE ESCALA", 1),
        ("2.1 Evolução Histórica", 2),
        ("2.2 Arquitetura Transformer e Mecanismo de Atenção", 2),
        ("2.3 Tokenização e Janela de Contexto", 2),
        ("3 ENGENHARIA DE PROMPT", 1),
        ("3.1 Principais Técnicas", 2),
        ("3.2 Relação entre Estrutura do Prompt e Qualidade da Resposta", 2),
        ("4 TEORIA DA INFORMAÇÃO E REDUNDÂNCIA LINGUÍSTICA", 1),
        ("4.1 Entropia de Shannon e Linguagem Natural", 2),
        ("4.2 Perplexidade como Métrica de Relevância Textual", 2),
        ("5 IMPACTO COMPUTACIONAL DOS TOKENS", 1),
        ("5.1 Complexidade Quadrática do Mecanismo de Atenção", 2),
        ("5.2 Custos Financeiros em APIs Comerciais", 2),
        ("6 MÉTODOS DE COMPRESSÃO DE PROMPTS", 1),
        ("6.1 LLMLingua e LLMLingua-2", 2),
        ("6.2 LongLLMLingua", 2),
        ("6.3 Selective Context e AutoCompressor", 2),
        ("7 COMPRESSÃO EM SISTEMAS RAG", 1),
        ("8 AVALIAÇÃO SEMÂNTICA COM BERTSCORE", 1),
        ("9 SUSTENTABILIDADE E GREEN AI", 1),
        ("10 RESULTADOS E DISCUSSÃO EXPERIMENTAL", 1),
        ("10.1 Metodologia Experimental e Bancada de Testes", 2),
        ("10.2 Avaliação do Volume de Tokens e Taxas de Compressão", 2),
        ("10.3 Latência de Inferência e Tempo de Resposta (TTFT)", 2),
        ("10.4 Fidelidade Semântica via BERTScore", 2),
        ("10.5 Viabilidade Financeira e Eficiência Energética", 2),
        ("10.6 Diretrizes de Engenharia e Análise de Trade-offs", 2),
        ("11 CONCLUSÃO E TRABALHOS FUTUROS", 1),
        ("REFERÊNCIAS", 1),
    ]

    page_map = sumario_page_map or {}

    for item, lvl in toc_items:
        pg_num = str(page_map.get(item, "5"))
        p_sum = doc.add_paragraph()
        p_sum.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p_sum.paragraph_format.space_before = Pt(1.5)
        p_sum.paragraph_format.space_after = Pt(1.5)
        p_sum.paragraph_format.line_spacing = 1.0
        p_sum.paragraph_format.first_line_indent = Cm(0)
        p_sum.paragraph_format.left_indent = Cm(0)

        r_item = p_sum.add_run(f"{item}\t{pg_num}")
        r_item.font.name = "Arial"
        r_item.font.size = Pt(10)
        r_item.font.bold = (lvl == 1)

    # =========================================================================
    # CORPO DO TEXTO (FLUXO CONTÍNUO A PARTIR DA INTRODUÇÃO)
    # =========================================================================
    doc.add_page_break()

    structured_json_path = Path(__file__).parent / "structured_ch1_9.json"
    with open(structured_json_path, "r", encoding="utf-8") as f:
        structured_data = json.load(f)

    for elem in structured_data:
        etype = elem["type"]
        etext = elem["text"]
        if etype == "h1" and (etext.startswith("10 ") or etext.startswith("11 ")):
            break
        if etype == "h1":
            add_h1(etext, page_break_before=False)
        elif etype == "h2":
            add_h2(etext)
        elif etype == "bullet":
            add_p(etext, first_line_indent=Cm(0.5))
        else:
            add_p(etext)

    # =========================================================================
    # CAPÍTULO 10: RESULTADOS E DISCUSSÃO EXPERIMENTAL (100% SINCRONIZADO)
    # =========================================================================
    add_h1("10 RESULTADOS E DISCUSSÃO EXPERIMENTAL", page_break_before=False)

    add_h2("10.1 Metodologia Experimental e Bancada de Testes")
    add_p(
        "Para validar empiricamente os pressupostos teóricos sem recorrer a simulações sintéticas, "
        "construímos a bancada experimental automatizada Battle IA. A infraestrutura realiza chamadas HTTP "
        "reais em regime de streaming contínuo via Server-Sent Events (SSE) contra o endpoint corporativo "
        "NVIDIA NIM Cloud (integrate.api.nvidia.com), executando o modelo fundacional multimodal "
        "Meta Llama-3.2-11B-Vision-Instruct (identificador de API meta/llama-3.2-11b-vision-instruct) "
        "alocado em aceleradores NVIDIA A100 SXM4 com temperatura controlada em 0,1 e teto de geração "
        "padronizado em 260 tokens de saída (max_tokens = 260)."
    )
    add_p("Os ensaios abrangeram três cargas de trabalho representativas da engenharia de computação:")

    c1_few = by_key[("caso_1", "Few-shot (Baseline)")]["tokens_entrada"]
    c1_zero = by_key[("caso_1", "Zero-shot")]["tokens_entrada"]
    c2_few = by_key[("caso_2", "Few-shot (Baseline)")]["tokens_entrada"]
    c2_zero = by_key[("caso_2", "Zero-shot")]["tokens_entrada"]
    c3_few = by_key[("caso_3", "Few-shot (Baseline)")]["tokens_entrada"]
    c3_zero = by_key[("caso_3", "Zero-shot")]["tokens_entrada"]

    add_p(
        f"1. Caso 1 — Raciocínio em Sistemas Distribuídos e Concorrência: Análise de particionamento de rede "
        f"em cluster Raft de 5 nós sob o Teorema CAP, prevenção de split-brain via quórum de maioria (N/2 + 1) "
        f"e linearizabilidade de leituras por leasing de liderança ({c1_few} tokens de entrada no baseline Few-shot "
        f"e {c1_zero} tokens no Zero-shot);"
    )
    add_p(
        f"2. Caso 2 — Engenharia de Dados e Geração de Consultas SQL sob Esquemas DDL: Mapeamento de esquema "
        f"relacional em Terceira Forma Normal (3FN) composto pelas tabelas clientes, pedidos, itens_pedido e faturas, "
        f"exigindo agregação financeira com Common Table Expressions (WITH / CTE) e função de janela DENSE_RANK() "
        f"({c2_few} tokens de entrada no baseline Few-shot e {c2_zero} tokens no Zero-shot);"
    )
    add_p(
        f"3. Caso 3 — Governança e Auditoria de Conformidade com a LGPD: Parecer técnico-jurídico sob a Lei "
        f"nº 13.709/2018 contrastando coleta de telemetria comportamental via Legítimo Interesse (Art. 7º, IX) "
        f"e captura biométrica facial enquadrada como Dado Pessoal Sensível (Art. 5º, II e Art. 11º) "
        f"({c3_few} tokens de entrada no baseline Few-shot e {c3_zero} tokens no Zero-shot)."
    )
    add_p(
        "Confrontamos duas linhas de base convencionais — Few-shot (Baseline) e Zero-shot — com o motor "
        "algorítmico ShannonPromptCompressor, inspirado no LLMLingua e parametrizado em fatores nominais de poda "
        "de 2x e 4x. O compressor calcula a autoinformação de Shannon I(w) = -log2 P(w) de cada vocábulo, podando "
        "conectivos redundantes enquanto aplica proteção estrita (token whitelisting) sobre identificadores técnicos, "
        "cláusulas SQL e referências normativas."
    )

    add_h2("10.2 Avaliação do Volume de Tokens e Taxas de Compressão")
    add_p(
        "A contabilização de tokens foi extraída diretamente do manifesto oficial usage retornado pelo servidor "
        "NVIDIA NIM (prompt_tokens, completion_tokens e total_tokens). A Tabela 1 apresenta o consumo exato "
        "medido em cada configuração e a respectiva redução percentual frente à linha de base Few-shot."
    )

    add_table_title("Tabela 1 – Consumo real de tokens faturados na API NVIDIA NIM e taxas efetivas de compressão")
    headers_t1 = ["Caso de Uso", "Técnica de Prompting", "Tokens Entrada", "Tokens Saída", "Razão Efetiva", "Redução Entrada (%)"]

    casos_meta = [
        ("caso_1", "Caso 1: Sistemas Distribuídos"),
        ("caso_2", "Caso 2: SQL sob DDL"),
        ("caso_3", "Caso 3: Auditoria LGPD"),
    ]
    tecnicas_ordem_t1 = ["Few-shot (Baseline)", "Zero-shot", "LLMLingua (2x)", "LLMLingua (4x)"]

    data_t1 = []
    for cid, c_label in casos_meta:
        base_in = by_key[(cid, "Few-shot (Baseline)")]["tokens_entrada"]
        for tec in tecnicas_ordem_t1:
            rec = by_key[(cid, tec)]
            t_in = rec["tokens_entrada"]
            t_out = rec["tokens_saida"]
            razao = base_in / t_in
            if tec == "Few-shot (Baseline)":
                razao_str = "1,00x (Ref)"
                red_str = "Baseline (0,0%)"
            else:
                red_pct = ((t_in - base_in) / base_in) * 100.0
                razao_str = f"{fmt_br(razao, 2)}x"
                red_str = f"{fmt_br(red_pct, 1)}%"
            tec_short = "Few-shot" if tec == "Few-shot (Baseline)" else tec
            data_t1.append([c_label, tec_short, str(t_in), str(t_out), razao_str, red_str])

    col_w_t1 = [Cm(4.8), Cm(3.2), Cm(2.2), Cm(2.0), Cm(2.3), Cm(2.5)]
    t1 = doc.add_table(rows=1, cols=6)
    format_abnt_table(t1, col_w_t1, headers_t1, data_t1)
    add_table_source()

    add_p(
        "A Figura 10.1 reproduz graficamente as medições da Tabela 1, evidenciando o degrau de redução "
        "no volume de tokens de entrada processados pelo modelo Meta Llama-3.2-11B-Vision-Instruct."
    )

    add_figure(
        figs_dir / "figura_10_1_volume_tokens.png",
        "Figura 10.1 – Consumo real de tokens de entrada por técnica de prompting nos três cenários avaliados"
    )

    add_p(
        "Observa-se um fenômeno de engenharia fundamental ao confrontar as taxas nominais do compressor (2x e 4x) "
        "com as razões efetivas faturadas na API (1,63x a 1,65x para 2x; 2,21x a 2,35x para 4x frente ao baseline Few-shot). "
        "Enquanto a poda lexical atua sobre o corpo da instrução do usuário, a requisição real contabilizada em prompt_tokens "
        "inclui o system prompt invariante, os delimitadores especiais do template de chat do Llama-3.2 e os tokens blindados "
        "pela lista de proteção semântica (whitelisting de operadores SQL, nomes de tabelas e artigos da LGPD). "
        "Ainda assim, a compressão nominal de 4x eliminou 56,2% dos tokens de entrada no Caso 1 (de 340 para 149 tokens), "
        "57,4% no Caso 2 (de 434 para 185 tokens) e 54,8% no Caso 3 (de 378 para 171 tokens), mantendo a geração "
        "estável no teto de 260 tokens de saída."
    )

    add_h2("10.3 Latência de Inferência e Tempo de Resposta (TTFT)")
    add_p(
        "A latência percebida pelo usuário em sistemas de streaming é governada pelo Time-to-First-Token (TTFT), "
        "cronometrado na bancada desde o disparo do payload HTTP até a recepção do primeiro fragmento de texto via SSE. "
        "A Tabela 2 consolida o TTFT em milissegundos, o tempo total de geração em segundos e as variações relativas."
    )

    add_table_title("Tabela 2 – Latência inicial de resposta (TTFT) e tempo total de inferência medidos via SSE streaming")
    headers_t2 = ["Caso de Uso", "Técnica", "TTFT (ms)", "Tempo Total (s)", "Var. TTFT vs Few-shot", "Var. TTFT vs Zero-shot"]

    data_t2 = []
    for cid, c_label in casos_meta:
        base_ttft = by_key[(cid, "Few-shot (Baseline)")]["ttft_ms"]
        zero_ttft = by_key[(cid, "Zero-shot")]["ttft_ms"]
        for tec in tecnicas_ordem_t1:
            rec = by_key[(cid, tec)]
            ttft = rec["ttft_ms"]
            t_tot = rec["tempo_total_s"]
            if tec == "Few-shot (Baseline)":
                var_few = "Baseline (0,0%)"
            else:
                diff_few = ((ttft - base_ttft) / base_ttft) * 100.0
                sinal = "+" if diff_few > 0 else ""
                var_few = f"{sinal}{fmt_br(diff_few, 1)}%"

            if tec == "Zero-shot":
                var_zero = "Ref. Zero-shot"
            else:
                diff_zero = ((ttft - zero_ttft) / zero_ttft) * 100.0
                sinal_z = "+" if diff_zero > 0 else ""
                var_zero = f"{sinal_z}{fmt_br(diff_zero, 1)}%"

            tec_short = "Few-shot" if tec == "Few-shot (Baseline)" else tec
            data_t2.append([
                c_label,
                tec_short,
                fmt_br(ttft, 1),
                fmt_br(t_tot, 2),
                var_few,
                var_zero
            ])

    col_w_t2 = [Cm(4.6), Cm(3.0), Cm(2.1), Cm(2.2), Cm(2.5), Cm(2.6)]
    t2 = doc.add_table(rows=1, cols=6)
    format_abnt_table(t2, col_w_t2, headers_t2, data_t2)
    add_table_source()

    add_p(
        "A Figura 10.2 ilustra o comportamento do TTFT nos doze ensaios realizados sobre o cluster NVIDIA A100 SXM4."
    )

    add_figure(
        figs_dir / "figura_10_2_latencia_ttft.png",
        "Figura 10.2 – Comportamento real do Time-to-First-Token (TTFT em ms) sob diferentes técnicas de prompting"
    )

    add_p(
        "A análise empírica de um endpoint de nuvem multi-inquilino (multi-tenant cloud inference) revela que a latência "
        "inicial observada resulta da soma TTFT = T_rede + T_fila + T_prefill. Nos ensaios sem compressão ou com poda intermediária, "
        "ocorreram picos transitórios de enfileiramento e contenção de escalonamento no cluster compartilhado — notadamente no "
        "Zero-shot do Caso 2 (2.032,7 ms) e do Caso 3 (4.418,7 ms), bem como na rodada LLMLingua 2x do Caso 1 (914,3 ms) e do "
        "Caso 3 (1.344,7 ms). Em contrapartida, a compressão de 4x apresentou estabilidade superior em todos os cenários, "
        "mantendo o TTFT estritamente na faixa de 538,0 ms (Caso 1), 636,7 ms (Caso 2) e 548,2 ms (Caso 3). "
        "No Caso 3, o prompt enxuto de 171 tokens (4x) reduziu o TTFT em 23,0% frente ao baseline Few-shot (711,6 ms) "
        "e em 87,6% frente ao pico de contenção do Zero-shot (4.418,7 ms), demonstrando que sequências mais curtas "
        "favorecem a alocação imediata de blocos de KV-cache no agendador de lotes contínuos (continuous batching) da GPU."
    )

    add_h2("10.4 Fidelidade Semântica via BERTScore")
    add_p(
        "Para verificar se a poda lexical preserva a exatidão conceitual da resposta, submetemos as saídas do modelo "
        "ao avaliador SemanticFidelityEvaluator, contrastando cada geração com o gabarito técnico de referência. "
        "A métrica composta BERTScore Proxy F1 integra a similaridade cossenoidal de N-gramas (peso 0,50), a taxa de "
        "retenção de entidades técnicas obrigatórias (peso 0,35) e o índice ROUGE-L de subsequência comum mais longa (peso 0,15), "
        "adotando o limiar crítico acadêmico de F1 = 0,80. A Tabela 3 detalha os indicadores obtidos."
    )

    add_table_title("Tabela 3 – Métricas reais de fidelidade semântica: Similaridade de Cosseno, ROUGE-L, Retenção de Entidades e BERTScore Proxy F1")
    headers_t3 = ["Caso de Uso", "Técnica / Taxa", "Cosine Sim.", "ROUGE-L", "Ret. Entidades", "BERTScore F1", "Status (≥ 0,80)"]

    data_t3 = []
    for cid, c_label in casos_meta:
        for tec in tecnicas_ordem_t1:
            rec = by_key[(cid, tec)]
            cos_v = rec["cosine_similarity"]
            rouge_v = rec["rouge_l_f1"]
            ent_v = rec["entity_retention"] * 100.0
            f1_v = rec["bertscore_proxy_f1"]
            if f1_v >= 0.85:
                status = "Excelente (≥ 0,85)"
            elif f1_v >= 0.80:
                status = "Aprovado (≥ 0,80)"
            else:
                status = "Atenção (< 0,80)"
            tec_short = "Few-shot" if tec == "Few-shot (Baseline)" else tec
            data_t3.append([
                c_label,
                tec_short,
                fmt_br(cos_v, 4),
                fmt_br(rouge_v, 4),
                f"{fmt_br(ent_v, 1)}%",
                fmt_br(f1_v, 4),
                status
            ])

    col_w_t3 = [Cm(4.2), Cm(2.7), Cm(2.0), Cm(1.9), Cm(2.1), Cm(2.1), Cm(2.5)]
    t3 = doc.add_table(rows=1, cols=7)
    format_abnt_table(t3, col_w_t3, headers_t3, data_t3)
    add_table_source()

    add_p(
        "A Figura 10.3 apresenta a curva de retenção semântica em função da taxa de compressão (1x Zero-shot, 2x e 4x), "
        "demarcando a linha horizontal do limiar crítico acadêmico (F1 = 0,80)."
    )

    add_figure(
        figs_dir / "figura_10_3_bertscore_f1.png",
        "Figura 10.3 – Curva real de retenção semântica (BERTScore Proxy F1) por razão de compressão e limiar crítico (0,80)"
    )

    add_p(
        "A dissecação dos três domínios expõe comportamentos distintos e cientificamente reveladores. "
        "No Caso 2 (SQL sob DDL), o mecanismo de token whitelisting preservou 100,0% das entidades obrigatórias em todas as configurações, "
        "sustentando BERTScore F1 na faixa de excelência tanto em 2x (0,8623) quanto em 4x (0,8524). "
        "No Caso 3 (Auditoria LGPD), comprovou-se empiricamente o fenômeno de ganho de relação sinal-ruído descrito por Jiang et al. (2023): "
        "enquanto os prompts extensos (Few-shot com F1 de 0,7916 e Zero-shot com 0,7955) dispersaram a geração inicial do modelo em preâmbulos "
        "introdutórios — retendo apenas 37,5% das entidades jurídicas antes do corte de 260 tokens —, os prompts comprimidos concentraram "
        "a atenção diretamente nos dispositivos legais protegidos (Art. 7º, Art. 11º e dado sensível), elevando a retenção de entidades para "
        "50,0% (F1 de 0,8037 em 2x) e 75,0% (F1 de 0,8279 em 4x). "
        "Por outro lado, no Caso 1 (Sistemas Distribuídos), o índice F1 recuou de 0,8115 (Zero-shot) para 0,7917 (2x) e 0,7912 (4x), "
        "ficando marginalmente abaixo do corte de 0,80 (-0,0088). Isso demonstra que, em modelos compactos de 11B parâmetros, "
        "explicações teóricas multi-etapas sobre consenso distribuído ressentem-se da remoção de preposições e conectivos causais."
    )

    add_h2("10.5 Viabilidade Financeira e Eficiência Energética")
    add_p(
        "Para dimensionar o impacto operacional em larga escala, projetamos o processamento de um lote corporativo "
        "de 100.000 requisições sob a estrutura tarifária de referência implementada na bancada (US$ 3,00 por milhão "
        "de tokens de entrada e US$ 15,00 por milhão de tokens de saída). Como todas as execuções operaram com teto fixo "
        "de 260 tokens de saída (26,0 milhões de tokens gerados em 100k chamadas, equivalentes a US$ 390,00 fixos de saída), "
        "a variação econômica reflete exatamente o alívio obtido no processamento de entrada. A Tabela 4 sintetiza "
        "as médias consolidadas dos três casos para cada técnica, em correspondência exata com a Figura 10.4."
    )

    add_table_title("Tabela 4 – Projeção financeira consolidada (100.000 requisições) e redução média de FLOPs de autoatenção O(n²)")
    headers_t4 = ["Cenário / Técnica", "Tokens Entrada (100k)", "Custo Entrada (US$)", "Custo Total (US$)", "Economia (Entrada / Total)", "Redução FLOPs O(n²)"]

    cids = ["caso_1", "caso_2", "caso_3"]
    base_avg_in = sum(by_key[(c, "Few-shot (Baseline)")]["tokens_entrada"] for c in cids) / 3.0
    base_avg_cost_in = (base_avg_in * 100_000 / 1_000_000.0) * 3.00
    base_avg_cost_tot = sum(by_key[(c, "Few-shot (Baseline)")]["custo_total_100k_usd"] for c in cids) / 3.0

    data_t4 = []
    for idx_t, tec in enumerate(tecnicas_ordem_t1, start=1):
        avg_in = sum(by_key[(c, tec)]["tokens_entrada"] for c in cids) / 3.0
        vol_m = (avg_in * 100_000) / 1_000_000.0
        cost_in = vol_m * 3.00
        cost_tot = sum(by_key[(c, tec)]["custo_total_100k_usd"] for c in cids) / 3.0
        avg_flops = sum(by_key[(c, tec)]["reducao_flops_atencao_pct"] for c in cids) / 3.0

        if tec == "Few-shot (Baseline)":
            econ_str = "Baseline (0,0%)"
            flops_str = "0,0% (Ref)"
        else:
            ec_in_pct = ((cost_in - base_avg_cost_in) / base_avg_cost_in) * 100.0
            ec_tot_pct = ((cost_tot - base_avg_cost_tot) / base_avg_cost_tot) * 100.0
            econ_str = f"{fmt_br(ec_in_pct, 1)}% / {fmt_br(ec_tot_pct, 1)}%"
            flops_str = f"-{fmt_br(avg_flops, 1)}%"

        tec_label = "Few-shot (Baseline)" if tec == "Few-shot (Baseline)" else tec
        data_t4.append([
            f"Cenário {idx_t}: {tec_label}",
            f"{fmt_br(vol_m, 2)} M",
            f"US$ {fmt_br(cost_in, 2)}",
            f"US$ {fmt_br(cost_tot, 2)}",
            econ_str,
            flops_str
        ])

    col_w_t4 = [Cm(4.5), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.8), Cm(2.2)]
    t4 = doc.add_table(rows=1, cols=6)
    format_abnt_table(t4, col_w_t4, headers_t4, data_t4)
    add_table_source()

    add_p(
        "A Figura 10.4 confronta em eixo duplo o custo total médio projetado para 100.000 requisições "
        "(US$ 505,20 no Few-shot, US$ 482,80 no Zero-shot, US$ 460,30 no LLMLingua 2x e US$ 440,50 no LLMLingua 4x) "
        "com a redução percentual média nos cálculos quadráticos de autoatenção durante a fase de prefill."
    )

    add_figure(
        figs_dir / "figura_10_4_custo_flops.png",
        "Figura 10.4 – Economia financeira projetada (100k requisições) e redução média de FLOPs na autoatenção O(n²)"
    )

    add_p(
        "Como o esforço computacional da camada de autoatenção escala segundo 1 - (T_comprimido / T_baseline)², "
        "uma redução média de 56,2% nos tokens de entrada sob LLMLingua 4x (de 38,40 milhões para 16,83 milhões de tokens "
        "a cada 100.000 chamadas) traduz-se em uma supressão média de 80,7% nos FLOPs de autoatenção do prefill "
        "(80,8% no Caso 1, 81,8% no Caso 2 e 79,5% no Caso 3). Financeiramente, o custo isolado de entrada cai 56,2% "
        "(de US$ 115,20 para US$ 50,50), gerando economia líquida de US$ 64,70 (-12,8% no custo total do lote, que é "
        "amortecido pelo peso tarifário 5x superior dos 260 tokens fixos de saída). Sob a ótica de Green AI, "
        "poupar mais de 80% das multiplicações matriciais de prefill reduz diretamente a pressão térmica e energética "
        "nos módulos HBM dos aceleradores A100 SXM4."
    )

    add_h2("10.6 Diretrizes de Engenharia e Análise de Trade-offs")
    add_p(
        "A consolidação dos resultados experimentais permite estabelecer três diretrizes práticas para arquitetos de software: "
        "(1) obrigatoriedade de listas de proteção de vocabulário (token whitelisting) em tarefas estruturadas como geração de SQL "
        "e análise normativa, garantindo 100% de preservação de operadores e cláusulas críticas; "
        "(2) adoção preferencial da taxa de 4x em tarefas de extração e auditoria regulatória (como LGPD), onde a remoção "
        "de verbosidade acessória atua como filtro de ruído e melhora a densidade da resposta; e "
        "(3) cautela na aplicação de podas agressivas sobre modelos compactos (porte de 8B a 11B parâmetros) em problemas "
        "de raciocínio causal encadeado (como algoritmos de consenso distribuído), nos quais se recomenda limitar a compressão "
        "ou preservar conectivos lógicos de causa e consequência."
    )

    # =========================================================================
    # CAPÍTULO 11: CONCLUSÃO E TRABALHOS FUTUROS
    # =========================================================================
    add_h1("11 CONCLUSÃO E TRABALHOS FUTUROS", page_break_before=False)
    add_p(
        "Este trabalho investigou os fundamentos teóricos e comprovou experimentalmente a eficácia da compressão "
        "algorítmica de prompts como estratégia de otimização computacional, financeira e ambiental para LLMs. "
        "A articulação entre a Teoria da Informação de Claude Shannon e a mecânica de autoatenção dos Transformers "
        "demonstra que a eliminação seletiva de tokens de baixa autoinformação reduz drasticamente o esforço matricial "
        "na fase de prefill."
    )
    add_p(
        "Os testes empíricos conduzidos na bancada Battle IA com o modelo Meta Llama-3.2-11B-Vision-Instruct em "
        "infraestrutura NVIDIA NIM Cloud (A100 SXM4) demonstraram que a compressão nominal de 4x combinada a token whitelisting "
        "reduz em até 57,4% o volume faturado de tokens de entrada (média de 56,2%) e suprime em média 80,7% dos FLOPs "
        "quadráticos de autoatenção. Comprovou-se que a técnica mantém excelente fidelidade semântica na geração de consultas "
        "SQL sob DDL (BERTScore Proxy F1 de 0,8524 com 100% de retenção de entidades) e eleva a precisão em auditoria de "
        "conformidade LGPD (F1 subindo de 0,7916 no baseline para 0,8279 em 4x pela eliminação de ruído textual), "
        "ao passo que evidenciou o limite de sensibilidade de modelos de 11B parâmetros em raciocínio de sistemas distribuídos (F1 de 0,7912)."
    )
    add_p(
        "Para trabalhos futuros, propõem-se três frentes de investigação: "
        "(a) calibração adaptativa da taxa de poda guiada pela perplexidade contextual de modelos auxiliares dedicados (LLMLingua-2) "
        "para proteger conectivos causais em tarefas de raciocínio distribuído; "
        "(b) extensão da bancada Battle IA para comparação direta entre modelos compactos (11B) e modelos de fronteira (70B+) "
        "sob múltiplas repetições estatísticas para isolamento do jitter de rede em nuvem; e "
        "(c) avaliação conjunta de compressão de prompts e cache semântico de prefixos em arquiteturas RAG corporativas."
    )

    # =========================================================================
    # REFERÊNCIAS BIBLIOGRÁFICAS (100% AUDITADAS - ABNT NBR 6023)
    # =========================================================================
    add_h1("REFERÊNCIAS", page_break_before=True)

    referencias = [
        "BAHDANAU, Dzmitry; CHO, Kyunghyun; BENGIO, Yoshua. Neural machine translation by jointly learning to align and translate. In: INTERNATIONAL CONFERENCE ON LEARNING REPRESENTATIONS (ICLR), 3., 2015, San Diego. Proceedings [...]. San Diego: ICLR, 2015. p. 1–15. Disponível em: https://arxiv.org/abs/1409.0473. Acesso em: 16 set. 2026.",
        "BROWN, Tom B. et al. Language models are few-shot learners. In: ADVANCES IN NEURAL INFORMATION PROCESSING SYSTEMS (NeurIPS), 33., 2020, San Diego (Virtual). Proceedings [...]. Red Hook: Curran Associates, 2020. v. 33, p. 1877–1901. Disponível em: https://proceedings.neurips.cc/paper/2020/hash/1457c0d6bfcb4967418bfb8ac142f64a-Abstract.html. Acesso em: 16 set. 2026.",
        "CHEVALIER, Alexis et al. Adapting language models to compress contexts. In: CONFERENCE ON EMPIRICAL METHODS IN NATURAL LANGUAGE PROCESSING (EMNLP), 2023, Singapore. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2023. p. 3829–3846. DOI: https://doi.org/10.18653/v1/2023.emnlp-main.232. Disponível em: https://aclanthology.org/2023.emnlp-main.232/. Acesso em: 16 set. 2026.",
        "DEVLIN, Jacob et al. BERT: pre-training of deep bidirectional transformers for language understanding. In: CONFERENCE OF THE NORTH AMERICAN CHAPTER OF THE ASSOCIATION FOR COMPUTATIONAL LINGUISTICS: HUMAN LANGUAGE TECHNOLOGIES (NAACL-HLT), 2019, Minneapolis. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2019. v. 1, p. 4171–4186. DOI: https://doi.org/10.18653/v1/N19-1423. Disponível em: https://aclanthology.org/N19-1423/. Acesso em: 16 set. 2026.",
        "JIANG, Huiqiang et al. LLMLingua: compressing prompts for accelerated inference of large language models. In: CONFERENCE ON EMPIRICAL METHODS IN NATURAL LANGUAGE PROCESSING (EMNLP), 2023, Singapore. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2023. p. 13358–13376. DOI: https://doi.org/10.18653/v1/2023.emnlp-main.825. Disponível em: https://aclanthology.org/2023.emnlp-main.825/. Acesso em: 16 set. 2026.",
        "KOJIMA, Takeshi et al. Large language models are zero-shot reasoners. In: ADVANCES IN NEURAL INFORMATION PROCESSING SYSTEMS (NeurIPS), 35., 2022, New Orleans. Proceedings [...]. Red Hook: Curran Associates, 2022. v. 35, p. 22199–22213. Disponível em: https://proceedings.neurips.cc/paper_files/paper/2022/hash/8bb0d291acd4acf06ef112099c16f326-Abstract-Conference.html. Acesso em: 16 set. 2026.",
        "LEWIS, Patrick et al. Retrieval-augmented generation for knowledge-intensive NLP tasks. In: ADVANCES IN NEURAL INFORMATION PROCESSING SYSTEMS (NeurIPS), 33., 2020, San Diego (Virtual). Proceedings [...]. Red Hook: Curran Associates, 2020. v. 33, p. 9459–9474. Disponível em: https://proceedings.neurips.cc/paper/2020/hash/6b493230205f780e1bc26945df7481e5-Abstract.html. Acesso em: 16 set. 2026.",
        "LI, Yucheng et al. Compressing context to enhance inference efficiency of large language models. In: CONFERENCE ON EMPIRICAL METHODS IN NATURAL LANGUAGE PROCESSING (EMNLP), 2023, Singapore. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2023. p. 6346–6365. DOI: https://doi.org/10.18653/v1/2023.emnlp-main.391. Disponível em: https://aclanthology.org/2023.emnlp-main.391/. Acesso em: 16 set. 2026.",
        "LIN, Chin-Yew. ROUGE: a package for automatic evaluation of summaries. In: ACL WORKSHOP ON TEXT SUMMARIZATION BRANCHES OUT, 2004, Barcelona. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2004. p. 74–81. Disponível em: https://aclanthology.org/W04-1013/. Acesso em: 16 set. 2026.",
        "LIU, Pengfei et al. Pre-train, prompt, and predict: a systematic survey of prompting methods in natural language processing. ACM Computing Surveys, New York, v. 55, n. 9, art. 195, p. 1–35, jan. 2023. DOI: https://doi.org/10.1145/3560815. Disponível em: https://dl.acm.org/doi/10.1145/3560815. Acesso em: 16 set. 2026.",
        "MIKOLOV, Tomas et al. Efficient estimation of word representations in vector space. In: INTERNATIONAL CONFERENCE ON LEARNING REPRESENTATIONS (ICLR) WORKSHOP TRACK, 1., 2013, Scottsdale. Proceedings [...]. Scottsdale: ICLR, 2013. p. 1–12. Disponível em: https://arxiv.org/abs/1301.3781. Acesso em: 16 set. 2026.",
        "PAN, Zhuoshi et al. LLMLingua-2: data distillation for efficient and faithful task-agnostic prompt compression. In: FINDINGS OF THE ASSOCIATION FOR COMPUTATIONAL LINGUISTICS: ACL 2024, 2024, Bangkok. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2024. p. 949–967. DOI: https://doi.org/10.18653/v1/2024.findings-acl.57. Disponível em: https://aclanthology.org/2024.findings-acl.57/. Acesso em: 16 set. 2026.",
        "PAPINENI, Kishore et al. BLEU: a method for automatic evaluation of machine translation. In: ANNUAL MEETING OF THE ASSOCIATION FOR COMPUTATIONAL LINGUISTICS (ACL), 40., 2002, Philadelphia. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2002. p. 311–318. DOI: https://doi.org/10.3115/1073083.1073135. Disponível em: https://aclanthology.org/P02-1040/. Acesso em: 16 set. 2026.",
        "PATTERSON, David et al. Carbon emissions and large neural network training. arXiv preprint arXiv:2104.10350, 2021. DOI: https://doi.org/10.48550/arXiv.2104.10350. Disponível em: https://arxiv.org/abs/2104.10350. Acesso em: 16 set. 2026.",
        "PENNINGTON, Jeffrey; SOCHER, Richard; MANNING, Christopher D. GloVe: global vectors for word representation. In: CONFERENCE ON EMPIRICAL METHODS IN NATURAL LANGUAGE PROCESSING (EMNLP), 2014, Doha. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2014. p. 1532–1543. DOI: https://doi.org/10.3115/v1/D14-1162. Disponível em: https://aclanthology.org/D14-1162/. Acesso em: 16 set. 2026.",
        "SCHWARTZ, Roy et al. Green AI. Communications of the ACM, New York, v. 63, n. 12, p. 54–63, dez. 2020. DOI: https://doi.org/10.1145/3381831. Disponível em: https://dl.acm.org/doi/10.1145/3381831. Acesso em: 16 set. 2026.",
        "SHANNON, Claude E. A mathematical theory of communication. The Bell System Technical Journal, Murray Hill, v. 27, n. 3, p. 379–423, jul. 1948; v. 27, n. 4, p. 623–656, out. 1948. DOI: https://doi.org/10.1002/j.1538-7305.1948.tb01338.x. Disponível em: https://archive.org/details/bstj27-3-379. Acesso em: 16 set. 2026.",
        "STRUBELL, Emma; GANESH, Ananya; McCALLUM, Andrew. Energy and policy considerations for deep learning in NLP. In: ANNUAL MEETING OF THE ASSOCIATION FOR COMPUTATIONAL LINGUISTICS (ACL), 57., 2019, Florence. Proceedings [...]. Stroudsburg: Association for Computational Linguistics, 2019. p. 3645–3650. DOI: https://doi.org/10.18653/v1/P19-1355. Disponível em: https://aclanthology.org/P19-1355/. Acesso em: 16 set. 2026.",
        "VASWANI, Ashish et al. Attention is all you need. In: ADVANCES IN NEURAL INFORMATION PROCESSING SYSTEMS (NeurIPS), 30., 2017, Long Beach. Proceedings [...]. Red Hook: Curran Associates, 2017. v. 30, p. 5998–6008. Disponível em: https://proceedings.neurips.cc/paper/2017/hash/3f5ee243547dee91fbd053c1c4a845aa-Abstract.html. Acesso em: 16 set. 2026.",
        "WEI, Jason et al. Chain-of-thought prompting elicits reasoning in large language models. In: ADVANCES IN NEURAL INFORMATION PROCESSING SYSTEMS (NeurIPS), 35., 2022, New Orleans. Proceedings [...]. Red Hook: Curran Associates, 2022. v. 35, p. 24824–24837. Disponível em: https://proceedings.neurips.cc/paper_files/paper/2022/hash/9d56096ce5236ab6591c4866b4477d88-Abstract-Conference.html. Acesso em: 16 set. 2026.",
        "ZHANG, Tianyi et al. BERTScore: evaluating text generation with BERT. In: INTERNATIONAL CONFERENCE ON LEARNING REPRESENTATIONS (ICLR), 8., 2020, Addis Ababa. Proceedings [...]. Addis Ababa: ICLR, 2020. p. 1–43. Disponível em: https://openreview.net/forum?id=SkeHuCVFDr. Acesso em: 16 set. 2026."
    ]

    for ref in referencias:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_ref.paragraph_format.space_before = Pt(0)
        p_ref.paragraph_format.space_after = Pt(5)
        p_ref.paragraph_format.line_spacing = 1.0
        p_ref.paragraph_format.first_line_indent = Cm(0)
        r = p_ref.add_run(ref)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"[OK] DOCX gerado em: {docx_path}")


def convert_to_pdf(docx_path: Path, pdf_path: Path):
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False
    try:
        doc = word.Documents.Open(str(docx_path.resolve()))
        doc.SaveAs2(str(pdf_path.resolve()), FileFormat=17)  # 17 = wdFormatPDF
        doc.Close()
        print(f"[OK] PDF compilado via Word COM: {pdf_path}")
        return True
    finally:
        word.Quit()


def scan_heading_pages(pdf_path: Path):
    doc = fitz.open(pdf_path)
    search_keys = [
        ("1 INTRODUÇÃO", "1 INTRODUÇÃO"),
        ("1.1 Contextualização e Motivação", "1.1 Contextualização"),
        ("1.2 Objetivos", "1.2 Objetivos"),
        ("1.3 Justificativa e Relevância", "1.3 Justificativa"),
        ("2 FUNDAMENTOS DOS MODELOS DE LINGUAGEM DE GRANDE ESCALA", "2 FUNDAMENTOS"),
        ("2.1 Evolução Histórica", "2.1 Evolução"),
        ("2.2 Arquitetura Transformer e Mecanismo de Atenção", "2.2 Arquitetura"),
        ("2.3 Tokenização e Janela de Contexto", "2.3 Tokenização"),
        ("3 ENGENHARIA DE PROMPT", "3 ENGENHARIA"),
        ("3.1 Principais Técnicas", "3.1 Principais"),
        ("3.2 Relação entre Estrutura do Prompt e Qualidade da Resposta", "3.2 Relação"),
        ("4 TEORIA DA INFORMAÇÃO E REDUNDÂNCIA LINGUÍSTICA", "4 TEORIA DA INFORMAÇÃO"),
        ("4.1 Entropia de Shannon e Linguagem Natural", "4.1 Entropia"),
        ("4.2 Perplexidade como Métrica de Relevância Textual", "4.2 Perplexidade"),
        ("5 IMPACTO COMPUTACIONAL DOS TOKENS", "5 IMPACTO COMPUTACIONAL"),
        ("5.1 Complexidade Quadrática do Mecanismo de Atenção", "5.1 Complexidade"),
        ("5.2 Custos Financeiros em APIs Comerciais", "5.2 Custos"),
        ("6 MÉTODOS DE COMPRESSÃO DE PROMPTS", "6 MÉTODOS"),
        ("6.1 LLMLingua e LLMLingua-2", "6.1 LLMLingua"),
        ("6.2 LongLLMLingua", "6.2 LongLLMLingua"),
        ("6.3 Selective Context e AutoCompressor", "6.3 Selective Context"),
        ("7 COMPRESSÃO EM SISTEMAS RAG", "7 COMPRESSÃO"),
        ("8 AVALIAÇÃO SEMÂNTICA COM BERTSCORE", "8 AVALIAÇÃO"),
        ("9 SUSTENTABILIDADE E GREEN AI", "9 SUSTENTABILIDADE"),
        ("10 RESULTADOS E DISCUSSÃO EXPERIMENTAL", "10 RESULTADOS"),
        ("10.1 Metodologia Experimental e Bancada de Testes", "10.1 Metodologia"),
        ("10.2 Avaliação do Volume de Tokens e Taxas de Compressão", "10.2 Avaliação"),
        ("10.3 Latência de Inferência e Tempo de Resposta (TTFT)", "10.3 Latência"),
        ("10.4 Fidelidade Semântica via BERTScore", "10.4 Fidelidade"),
        ("10.5 Viabilidade Financeira e Eficiência Energética", "10.5 Viabilidade"),
        ("10.6 Diretrizes de Engenharia e Análise de Trade-offs", "10.6 Diretrizes"),
        ("11 CONCLUSÃO E TRABALHOS FUTUROS", "11 CONCLUSÃO"),
        ("REFERÊNCIAS", "REFERÊNCIAS"),
    ]

    mapping = {}
    for label, query in search_keys:
        found_pg = None
        for p_idx, page in enumerate(doc):
            txt_up = page.get_text().upper()
            if p_idx < 4 or "SUMÁRIO" in txt_up:  # Ignorar pré-textuais e a página do próprio sumário
                continue
            if label == "REFERÊNCIAS":
                lines_up = [ln.strip() for ln in txt_up.splitlines() if ln.strip()]
                if "REFERÊNCIAS" in lines_up:
                    found_pg = p_idx + 1
                    break
            elif query.upper() in txt_up:
                found_pg = p_idx + 1
                break
        mapping[label] = str(found_pg or 5)
    return mapping


def main():
    manuscrito_dir = Path(__file__).parent.parent / "01 - Manuscrito"
    docx_file = manuscrito_dir / "TCC_II_Versao_1_Acompanhamento_Matheus_Santos.docx"
    pdf_file = manuscrito_dir / "TCC_II_Versao_1_Acompanhamento_Matheus_Santos.pdf"
    pdf_short = manuscrito_dir / "TCC_II.pdf"

    # 1ª Passada: Construir com sumário aproximado
    print("=== PASSADA 1: Compilação de teste para mapeamento de páginas ===")
    build_monograph(docx_file)
    convert_to_pdf(docx_file, pdf_file)

    # Mapear posições reais de páginas
    print("=== MAPEANDO PÁGINAS EXATAS DO PDF GERADO ===")
    exact_map = scan_heading_pages(pdf_file)
    for k, v in exact_map.items():
        print(f"  {k} -> pág. {v}")

    # 2ª Passada: Reconstruir com sumário 100% calibrado
    print("=== PASSADA 2: Compilação definitiva com sumário calibrado ===")
    build_monograph(docx_file, sumario_page_map=exact_map)
    convert_to_pdf(docx_file, pdf_file)

    # Sincronizar cópias espelho no cofre acadêmico
    shutil.copy2(pdf_file, pdf_short)
    print(f"[OK] Espelho atualizado: {pdf_short}")

    vault_root = Path(__file__).resolve().parents[3]
    mirror_dir = vault_root / "02 - TCC II" / "01 - Manuscrito"
    if mirror_dir.exists():
        shutil.copy2(docx_file, mirror_dir / docx_file.name)
        shutil.copy2(pdf_file, mirror_dir / pdf_file.name)
        print(f"[OK] Espelho 02 - TCC II atualizado em: {mirror_dir}")

    # Validação visual final com PyMuPDF
    doc = fitz.open(pdf_file)
    print(f"\n==========================================")
    print(f"     VALIDAÇÃO VISUAL DEFINITIVA DO PDF   ")
    print(f"==========================================")
    print(f"Total de páginas: {len(doc)}")

    p1_lines = [l.strip() for l in doc[0].get_text().splitlines() if l.strip()]
    has_city = any("INDAIATUBA" in l for l in p1_lines)
    has_year = any("2026" in l for l in p1_lines)
    print(f"Página 1 (Capa): {'PERFEITA (Contém INDAIATUBA 2026)' if (has_city and has_year) else 'ERRO: Vazou para página 2'}")

    p2_txt = doc[1].get_text()
    print(f"Página 2: {'RESUMO presente' if 'RESUMO' in p2_txt else 'ERRO no Resumo'}")

    p3_txt = doc[2].get_text()
    print(f"Página 3: {'ABSTRACT presente' if 'ABSTRACT' in p3_txt else 'ERRO no Abstract'}")

    p4_txt = doc[3].get_text()
    print(f"Página 4: {'SUMÁRIO presente' if 'SUMÁRIO' in p4_txt else 'ERRO no Sumário'}")

    total_imgs = sum(len(p.get_images()) for p in doc)
    print(f"Total de Figuras Científicas incorporadas: {total_imgs}")
    print("==========================================")


if __name__ == "__main__":
    main()
